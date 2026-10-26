"""Complex-slide import preserves page identity and fails without partial output."""
import io
from pathlib import Path
import shutil
import subprocess
import sys
import zipfile

import numpy as np
import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from ai_text_sharpener import ppt_raster, studio_project
from ai_text_sharpener.studio_project import StudioStore, export_vector_pptx, image_inputs
from ai_text_sharpener.typography import candidate_fonts


requires_renderer = pytest.mark.skipif(
    not ((shutil.which("libreoffice") or shutil.which("soffice"))
         and shutil.which("pdfinfo") and shutil.which("pdftoppm")),
    reason="LibreOffice and Poppler are required for the real PPT import test")


def available_tools(monkeypatch, missing=()):
    monkeypatch.setattr(ppt_raster.shutil, "which",
                        lambda name: None if name in missing else f"/test-tools/{name}")


@pytest.mark.parametrize("missing", [
    ("libreoffice", "soffice"), ("pdfinfo",), ("pdftoppm",),
])
def test_missing_renderer_components_have_actionable_error(monkeypatch, missing):
    available_tools(monkeypatch, missing)
    monkeypatch.setattr(ppt_raster, "_run", lambda *a, **kw: pytest.fail("must not run without tools"))
    with pytest.raises(ValueError, match="LibreOffice.*Poppler"):
        ppt_raster.render_pptx_pages(b"pptx", 1, [1])


def test_successful_converter_exit_without_pdf_is_failure(monkeypatch):
    available_tools(monkeypatch)
    paths = []

    def run(command, timeout, env=None):
        paths.append(Path(command[-1]).parent)
        return subprocess.CompletedProcess(command, 0, b"", b"")

    monkeypatch.setattr(ppt_raster, "_run", run)
    with pytest.raises(ValueError, match="未生成页面"):
        ppt_raster.render_pptx_pages(b"pptx", 1, [1])
    assert paths and all(not path.exists() for path in paths)


def fake_converter(monkeypatch, *, info=b"Pages: 5\n", omit_png=False):
    available_tools(monkeypatch)
    calls = []

    def run(command, timeout, env=None):
        calls.append(command)
        program = Path(command[0]).name
        if program in {"libreoffice", "soffice"}:
            source = Path(command[-1])
            assert source.read_bytes() == b"pptx"
            source.with_suffix(".pdf").write_bytes(b"test PDF")
        elif program == "pdfinfo":
            return subprocess.CompletedProcess(command, 0, info, b"")
        elif program == "pdftoppm" and not omit_png:
            number = int(command[command.index("-f") + 1])
            Image.new("RGB", (80, 45), (number, 20, 30)).save(Path(command[-1]).with_suffix(".png"))
        return subprocess.CompletedProcess(command, 0, b"", b"")

    monkeypatch.setattr(ppt_raster, "_run", run)
    return calls


@pytest.mark.parametrize("info", [b"Pages: 4\n", b"Title: no page count\n"])
def test_pdf_page_count_mismatch_stops_before_rasterization(monkeypatch, info):
    calls = fake_converter(monkeypatch, info=info)
    with pytest.raises(ValueError, match="页数不一致"):
        ppt_raster.render_pptx_pages(b"pptx", 5, [2])
    assert not any(Path(command[0]).name == "pdftoppm" for command in calls)


def test_selected_pages_keep_original_indices_and_survive_temp_cleanup(monkeypatch):
    calls = fake_converter(monkeypatch)
    images = ppt_raster.render_pptx_pages(b"pptx", 5, [5, 2])
    assert set(images) == {2, 5}
    assert images[2].size == (80, 45)
    assert images[2].getpixel((20, 20)) == (2, 20, 30)
    assert images[5].getpixel((20, 20)) == (5, 20, 30)
    for command in calls:
        if Path(command[0]).name == "pdftoppm":
            assert command[command.index("-f") + 1] == command[command.index("-l") + 1]
            assert not Path(command[-1]).parent.exists()


def test_missing_rendered_page_is_reported(monkeypatch):
    fake_converter(monkeypatch, omit_png=True)
    with pytest.raises(ValueError, match="第 2 页转换失败"):
        ppt_raster.render_pptx_pages(b"pptx", 5, [2])


def test_nonzero_converter_exit_is_reported():
    with pytest.raises(ValueError, match="转换失败"):
        ppt_raster._run([sys.executable, "-c", "raise SystemExit(7)"], timeout=5)


def test_converter_timeout_is_reported():
    with pytest.raises(ValueError, match="转换超时"):
        ppt_raster._run([sys.executable, "-c", "import time; time.sleep(10)"], timeout=.05)


def mixed_deck():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(12), Inches(6.75)
    first = Image.new("RGB", (640, 360), "#17496d")
    first.putpixel((0, 0), (255, 10, 20))
    last = Image.new("RGB", (640, 360), "#733755")
    last.putpixel((639, 359), (20, 240, 10))

    def slide_image(image):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        data = io.BytesIO()
        image.save(data, format="PNG")
        picture = slide.shapes.add_picture(io.BytesIO(data.getvalue()), 0, 0,
                                           deck.slide_width, deck.slide_height)
        return slide, picture

    def label(slide, text):
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(9), Inches(1.4))
        box.text = text
        box.text_frame.paragraphs[0].font.size = Pt(36)
        box.text_frame.paragraphs[0].font.name = "Liberation Sans"
        box.text_frame.paragraphs[0].font.color.rgb = RGBColor(10, 10, 10)

    slide_image(first)
    slide, _ = slide_image(Image.new("RGB", (640, 360), "#f0c040"))
    label(slide, "Native text on an image")
    hidden = deck.slides.add_slide(deck.slide_layouts[6])
    hidden._element.set("show", "0")
    hidden.background.fill.solid()
    hidden.background.fill.fore_color.rgb = RGBColor(32, 160, 64)
    label(hidden, "Hidden page stays third")
    split = Image.new("RGB", (640, 360), "#e02030")
    split.paste("#2060d0", (320, 0, 640, 360))
    _, picture = slide_image(split)
    picture.crop_left = .25
    picture.rotation = 12
    slide_image(last)
    data = io.BytesIO()
    deck.save(data)
    return data.getvalue(), first, last


@requires_renderer
def test_real_mixed_ppt_import_keeps_hidden_order_and_original_raster_pixels():
    data, first, last = mixed_deck()
    pages = image_inputs(data, "mixed.pptx")
    assert [name for name, _ in pages] == [f"第 {i} 页" for i in range(1, 6)]
    assert pages[0][1].size == first.size
    assert pages[0][1].tobytes() == first.tobytes()
    assert pages[4][1].size == last.size
    assert pages[4][1].tobytes() == last.tobytes()
    for _, image in pages[1:4]:
        assert abs(image.width / image.height - 16 / 9) < .01

    # The native text must be present, not just the underlying amber image.
    composed = np.asarray(pages[1][1].convert("RGB"))
    assert np.mean(np.max(np.abs(composed.astype(int) - (240, 192, 64)), axis=2) < 5) > .8
    assert np.mean(composed.max(axis=2) < 60) > .001
    hidden = np.asarray(pages[2][1].convert("RGB"))
    assert np.mean(np.max(np.abs(hidden.astype(int) - (32, 160, 64)), axis=2) < 5) > .8
    rotated = np.asarray(pages[3][1].convert("RGB"))
    assert np.mean((rotated[:, :, 0] > 180) & (rotated[:, :, 2] < 80)) > .1
    assert np.mean((rotated[:, :, 2] > 150) & (rotated[:, :, 0] < 80)) > .1
    assert np.mean(rotated.min(axis=2) > 245) > .01


@requires_renderer
def test_transparent_full_page_image_keeps_slide_background():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(32, 160, 64)
    source = Image.new("RGBA", (400, 300), (0, 0, 0, 0))
    source.paste((220, 30, 50, 255), (150, 100, 250, 200))
    raw = io.BytesIO()
    source.save(raw, format="PNG")
    slide.shapes.add_picture(io.BytesIO(raw.getvalue()), 0, 0, deck.slide_width, deck.slide_height)
    raw = io.BytesIO()
    deck.save(raw)
    image = image_inputs(raw.getvalue(), "transparent.pptx")[0][1]
    assert image.getpixel((image.width // 10, image.height // 10)) == (32, 160, 64)
    assert image.getpixel((image.width // 2, image.height // 2)) == (220, 30, 50)


@requires_renderer
def test_studio_vector_ppt_export_can_be_imported_again(tmp_path):
    fonts = candidate_fonts("Vector", 1)
    if not fonts:
        pytest.skip("A Latin font is required for vector text rendering")
    store = StudioStore(tmp_path)
    project = store.create("Vector round trip")
    store.add_images(project, [("Vector", Image.new("RGB", (640, 360), "#f0f0f0"))])
    page = project["pages"][0]
    page["regions"] = [{
        "id": "vector-text", "text": "Vector", "bbox": [40, 60, 260, 80],
        "x": 40, "y": 60, "font_id": fonts[0], "font_size": 64,
        "letter_spacing": 0, "color": "#dc1e32", "enabled": True, "erase_mode": "none",
    }]
    exported = export_vector_pptx(store.directory(project["id"]), project)
    with zipfile.ZipFile(exported) as package:
        assert b"svgBlip" in package.read("ppt/slides/slide1.xml")
        assert b"<path " in package.read("ppt/media/image_vector_1.svg")
    pages = image_inputs(exported.read_bytes(), exported.name)
    assert len(pages) == 1
    image = pages[0][1]
    pixels = np.asarray(image)
    assert np.mean((pixels[:, :, 0] > 180) & (pixels[:, :, 1] < 90)) > .005
    reference = Image.open(store.directory(project["id"]) / f"{page['id']}-result.png").convert("RGB")
    reduced = image.resize(reference.size, Image.Resampling.LANCZOS)
    assert np.abs(np.asarray(reduced).astype(float) - np.asarray(reference)).mean() < 3


def test_oversize_raster_is_rejected_without_starting_converter(monkeypatch):
    data, _, _ = mixed_deck()
    monkeypatch.setattr(studio_project, "MAX_PIXELS", 640 * 360 - 1)
    monkeypatch.setattr(studio_project, "render_pptx_pages",
                        lambda *a, **kw: pytest.fail("oversize input must not reach the converter"))
    with pytest.raises(ValueError, match="图片过大"):
        image_inputs(data, "oversize.pptx")


def test_invalid_or_empty_ppt_does_not_reach_converter(monkeypatch):
    monkeypatch.setattr(studio_project, "render_pptx_pages",
                        lambda *a, **kw: pytest.fail("invalid input must not reach the converter"))
    with pytest.raises(zipfile.BadZipFile):
        image_inputs(b"invalid zip", "invalid.pptx")
    raw = io.BytesIO()
    Presentation().save(raw)
    with pytest.raises(ValueError, match="没有页面"):
        image_inputs(raw.getvalue(), "empty.pptx")
