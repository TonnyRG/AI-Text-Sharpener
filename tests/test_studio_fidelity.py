"""Real rendering and persistence regression tests; no OCR download required."""
import io
import json
import threading
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

import numpy as np
import pytest
from PIL import Image, ImageFilter
from lxml import etree

from ai_text_sharpener.fidelity import crop_evidence, erase_regions, fit_region
from ai_text_sharpener.studio import create_server
from ai_text_sharpener.studio_project import StudioStore, extract_image_slides, export_vector_pptx, render_page
from ai_text_sharpener.style import sample_background_color
from ai_text_sharpener.erase import soft_erase
from ai_text_sharpener.detect import TextRegion
from ai_text_sharpener.style import RegionStyle
from ai_text_sharpener.typography import candidate_fonts, family_weight_fonts, font_catalog, outline_layout, render_mask


@pytest.fixture(scope="module")
def fonts():
    try:
        ids = candidate_fonts("Sample", 3)
    except FileNotFoundError:
        pytest.skip("fontconfig is required for Studio")
    if not ids:
        pytest.skip("No suitable fonts")
    return ids


def sample(fonts, text="Sample text", size=44):
    mask, layout = render_mask(fonts[0], text, size, 1.0)
    image = np.full((130, 700, 3), 240, np.uint8)
    h, w = mask.shape
    image[35:35+h, 40:40+w] = (240*(1-mask[...,None]) + 30*mask[...,None]).astype(np.uint8)
    image = np.array(Image.fromarray(image).filter(ImageFilter.GaussianBlur(.65)))
    region = {"id":"text-1", "text":text,"original_text":text,"bbox":[30,25,w+20,h+20],
              "x":40,"y":35,"font_id":fonts[0],"font_size":size,"letter_spacing":1.0,
              "color":"#1e1e1e","enabled":True,"erase_mode":"gradient","locked":False}
    return image, region


def test_fit_recovers_font_and_geometry(fonts):
    image, region = sample(fonts)
    result = fit_region(image, region, fonts)
    assert result["font_id"] == fonts[0]
    assert abs(result["font_size"] - 44) < 3
    assert abs(result["x"] - 40) < 3
    assert abs(result["y"] - 35) < 3
    assert result["score"] > .88


@pytest.mark.parametrize('decoration,color', [('bar', (165, 30, 65)), ('circle', (30, 30, 30))])
def test_fit_checks_text_extent_instead_of_including_decoration(fonts, decoration, color):
    import cv2
    from ai_text_sharpener.studio import refine_recognized_bounds
    image, region = sample(fonts)
    image = np.pad(image, ((0, 0), (40, 0), (0, 0)), constant_values=240)
    region['bbox'][0] = 15
    region['bbox'][2] += 65
    region['x'] += 40
    if decoration == 'bar':
        cv2.rectangle(image, (25, 35), (32, 65), color, -1)
    else:
        cv2.circle(image, (29, 50), 7, color, -1)
    original_bbox = region['bbox'][:]
    refine_recognized_bounds(image, region, [('Sample text', .99, [[76,25],[500,25],[500,85],[76,85]])])
    fitted = fit_region(image, region, fonts)
    assert fitted['bbox'][0] > 40
    assert abs(fitted['x'] - 80) < 3
    assert fitted['ocr_bbox'] == original_bbox
    for mode in ('gradient', 'inpaint'):
        erased = erase_regions(image, [dict(fitted, erase_mode=mode)])
        assert np.array_equal(erased[:, :40], image[:, :40])
        assert not np.array_equal(erased[35:65, 80:180], image[35:65, 80:180])


def test_fitting_preserves_recognized_isolated_letter(fonts):
    from ai_text_sharpener.studio import refine_recognized_bounds
    image, region = sample(fonts, text='I   Sample text')
    original_bbox=region['bbox'][:]
    # Colored text belongs to the recognizer support and must not be excluded.
    ink=image[35:75,40:50].min(axis=2)<150
    image[35:75,40:50][ink]=(165,30,65)
    refine_recognized_bounds(image,region,[('I   Sample text',.99,[[35,25],[500,25],[500,90],[35,90]])])
    assert region['bbox']==original_bbox
    fitted = fit_region(image, region, fonts)
    assert fitted['bbox'] == region['bbox']
    assert abs(fitted['x'] - 40) < 3


def test_incomplete_recognition_alignment_does_not_shrink_box(fonts):
    from ai_text_sharpener.studio import refine_recognized_bounds
    image,region=sample(fonts);before=region['bbox'][:]
    refine_recognized_bounds(image,region,[('text',.99,[[200,25],[300,25],[300,85],[200,85]])])
    assert region['bbox']==before


def test_detection_uses_recognizer_positions_to_exclude_adjacent_art(monkeypatch, fonts):
    from types import SimpleNamespace
    from ai_text_sharpener import studio
    image_path=Path(__file__).parent/'fixtures'/'studio'/'heading-with-adjacent-art.png'
    image=np.array(Image.open(image_path).convert('RGB'))
    text='衬线、中英文与数字'
    # Alignment recorded from RapidOCR on the unmodified generated test slide.
    limits=[(56,115),(120,175),(175,226),(226,281),(286,345),(350,409),(414,469),(469,524),(529,588)]
    words=tuple((char,.999,[[left,3],[right,3],[right,80],[left,80]]) for char,(left,right) in zip(text,limits))
    def recognize(path, **kwargs):
        assert kwargs['return_word_box'] is True
        return SimpleNamespace(boxes=np.array([[[3,3],[598,3],[598,80],[3,80]]]),txts=[text],scores=[.999],word_results=(words,))
    monkeypatch.setattr(studio,'ocr_engine',lambda:recognize)
    monkeypatch.setattr(studio,'candidate_fonts',lambda *args,**kwargs:fonts[:1])
    region=studio.detect_regions(image_path)[0]
    assert region['bbox'][0]>=55
    for mode in ('gradient','inpaint'):
        erased=erase_regions(image,[dict(region,erase_mode=mode)])
        assert np.array_equal(erased[:,:40],image[:,:40])
        assert not np.array_equal(erased[:,60:],image[:,60:])


def test_empty_text_erases_original_but_deleted_region_restores_it(tmp_path, fonts):
    from ai_text_sharpener.studio_project import compose_page
    import resvg_py
    image, region = sample(fonts)
    store=StudioStore(tmp_path);project=store.create('empty')
    store.add_images(project,[('slide',Image.fromarray(image))])
    page=project['pages'][0];region['text']='';page['regions']=[region]
    svg=compose_page(store.directory(project['id']),page)
    rendered=np.array(Image.open(io.BytesIO(bytes(resvg_py.svg_to_bytes(svg_string=svg)))).convert('RGB'))
    assert not np.array_equal(rendered,image)
    assert page['regions'][0]['ink_width']==0
    assert '<path ' not in svg
    page['regions']=[]
    svg=compose_page(store.directory(project['id']),page)
    restored=np.array(Image.open(io.BytesIO(bytes(resvg_py.svg_to_bytes(svg_string=svg)))).convert('RGB'))
    assert np.array_equal(restored,image)


def test_preview_renders_unsaved_edits_without_mutating_project(tmp_path, fonts):
    from ai_text_sharpener.studio_project import compose_page
    from copy import deepcopy
    server=create_server(tmp_path,0)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    image, region=sample(fonts)
    project=server.app.store.create('preview')
    server.app.store.add_images(project,[('slide',Image.fromarray(image))])
    page=project['pages'][0];page['regions']=[region];server.app.store.save(project)
    directory=server.app.store.directory(project['id'])
    before=(directory/'project.json').read_bytes()
    def preview(regions):
        body={'project_id':project['id'],'page_id':page['id'],'regions':regions}
        req=urllib.request.Request(f'http://127.0.0.1:{server.server_port}/api/preview',data=json.dumps(body).encode(),headers={'X-Studio-Token':server.app.token})
        with urllib.request.urlopen(req,timeout=10) as response:return json.load(response)
    try:
        changed=deepcopy(page);changed['regions'][0].update(text='Added text 123',x=135,y=60,bbox=[25,20,300,75])
        result=preview(changed['regions'])
        assert result['svg']==compose_page(directory,changed)
        assert 'translate(135 60)' in result['svg']
        assert result['regions'][0]['ink_width']>0
        assert (directory/'project.json').read_bytes()==before
        assert not list(directory.glob('*-result.*'))
        changed['regions'][0]['bbox'][2]=2
        with pytest.raises(urllib.error.HTTPError) as exc:preview(changed['regions'])
        assert exc.value.code==400
    finally:
        server.shutdown();server.server_close();thread.join();server.app.executor.shutdown()


def test_real_width_distinguishes_narrow_and_wide_letters(fonts):
    narrow = outline_layout(fonts[0], "iiiiiiii", 40., 0.)
    wide = outline_layout(fonts[0], "WWWWWWWW", 40., 0.)
    assert wide["width"] > narrow["width"] * 2


def test_font_priority_survives_localized_family_names(monkeypatch):
    from ai_text_sharpener import typography
    catalog = {
        "fallback": {"id":"fallback", "family":"FandolHei", "style":"Bold"},
        "yahei": {"id":"yahei", "family":"微软雅黑", "aliases":["微软雅黑", "Microsoft YaHei"], "style":"Bold"},
    }
    monkeypatch.setattr(typography, "font_catalog", lambda: catalog)
    monkeypatch.setattr(typography, "supports", lambda *_: True)
    assert typography.candidate_fonts("测试", 1) == ["yahei"]


def test_weight_refinement_finds_regular_omitted_from_initial_budget(monkeypatch):
    from ai_text_sharpener import fidelity
    family = {f['style']: f['id'] for f in font_catalog().values() if 'Source Han Sans CN' in f.get('aliases', [])}
    if not all(s in family for s in ['Normal','Regular','Light','Bold','Medium']):
        pytest.skip('Source Han Sans weight family not installed')
    selected = [family[s] for s in ['Normal','Light','Bold']]
    assert family['Regular'] in family_weight_fonts(selected, '研究生新生见面会')
    assert family['Medium'] in family_weight_fonts(selected, '研究生新生见面会')
    image, region = sample([family['Regular']], text='研究生新生见面会', size=44)
    monkeypatch.setattr(fidelity, 'candidate_fonts', lambda text: selected.copy())
    result = fidelity.fit_region(image, region)
    assert result['font_id'] == family['Regular']


def test_subpixel_stroke_increases_ink_without_changing_advances(fonts):
    plain, before = render_mask(fonts[0], 'Sample text', 44, 1)
    thick, after = render_mask(fonts[0], 'Sample text', 44, 1, .4)
    assert thick.sum() > plain.sum() * 1.03
    assert after['width'] == pytest.approx(before['width'] + .4)
    assert after['advance'] == before['advance']
    assert 'stroke="currentColor"' in after['body']


def test_stroke_is_validated_and_preserved_in_svg_and_ppt(tmp_path, fonts):
    from ai_text_sharpener.studio_project import validate_regions
    image, region = sample(fonts)
    store=StudioStore(tmp_path);project=store.create('stroke')
    store.add_images(project,[('slide',Image.fromarray(image))])
    page=project['pages'][0];page['regions']=[region]
    region['stroke_width']=float('nan')
    with pytest.raises(ValueError):validate_regions([region],700,130)
    region['stroke_width']=-.1
    with pytest.raises(ValueError):validate_regions([region],700,130)
    region['stroke_width']=.4
    exported=export_vector_pptx(store.directory(project['id']),project)
    with zipfile.ZipFile(exported) as archive:
        svg=archive.read('ppt/media/image_vector_1.svg').decode()
        assert 'stroke="currentColor"' in svg
        assert 'color="#1e1e1e"' in svg
    expected=outline_layout(region['font_id'],region['text'],44,1,.4)
    assert page['regions'][0]['ink_width']==round(expected['width'],2)


def test_black_background_is_valid():
    assert sample_background_color(np.zeros((60, 100, 3), np.uint8), (20,20,40,20)) == (0,0,0)


def test_legacy_feather_has_a_real_transition():
    source = np.zeros((120, 120, 3), np.uint8)
    region = TextRegion([[20,20],[100,20],[100,100],[20,100]],"x",1.)
    out = soft_erase(source,[(region,RegionStyle((0,0,0),(240,240,240),20))])
    assert 0 < out[20,60,0] < out[60,60,0]
    assert out[10,60,0] == 0


def test_erase_preserves_pixels_outside_selected_region(fonts):
    image, region = sample(fonts)
    result = erase_regions(image, [region])
    assert np.array_equal(result[:20], image[:20])
    assert np.array_equal(erase_regions(image,[dict(region,enabled=False)]), image)


def test_fit_blank_region_is_explicit_error(fonts):
    _,region=sample(fonts)
    with pytest.raises(ValueError, match="对比度"):
        fit_region(np.full((130,700,3),255,np.uint8),region,fonts)


def test_vector_export_retains_paths_and_png_fallback(tmp_path, fonts):
    image, region = sample(fonts)
    store=StudioStore(tmp_path)
    project=store.create("Test")
    store.add_images(project,[("slide",Image.fromarray(image))])
    page=project["pages"][0];page["regions"]=[region]
    path=export_vector_pptx(store.directory(project["id"]),project)
    with zipfile.ZipFile(path) as archive:
        svg=archive.read('ppt/media/image_vector_1.svg').decode()
        assert '<path ' in svg and '<text' not in svg
        assert 'data:image/png;base64,' in svg
        slide=etree.fromstring(archive.read('ppt/slides/slide1.xml'))
        assert len(slide.xpath('//*[local-name()="svgBlip"]'))==1
        assert len(slide.xpath('//*[local-name()="pic"]'))==1
        assert any(n.endswith('.png') for n in archive.namelist())
        rel=archive.read('ppt/slides/_rels/slide1.xml.rels').decode()
        assert 'image_vector_1.svg' in rel


def test_ppt_import_refuses_composed_slide():
    from pptx import Presentation
    from pptx.util import Inches
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);s.shapes.add_textbox(0,0,Inches(2),Inches(1)).text='Native text'
    raw=io.BytesIO();p.save(raw)
    with pytest.raises(ValueError,match="不是单张"):
        extract_image_slides(raw.getvalue())


def test_ppt_import_keeps_original_image_pixels():
    from pptx import Presentation
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6])
    im=Image.new('RGB',(400,300),(11,23,37));buf=io.BytesIO();im.save(buf,format='PNG');buf.seek(0)
    s.shapes.add_picture(buf,0,0,width=p.slide_width,height=p.slide_height)
    raw=io.BytesIO();p.save(raw)
    images=extract_image_slides(raw.getvalue())
    assert len(images)==1 and images[0][1].size==(400,300)
    assert images[0][1].getpixel((20,20))==(11,23,37)


def test_server_upload_save_revision_and_token(tmp_path, fonts):
    server=create_server(tmp_path,0)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    base=f'http://127.0.0.1:{server.server_port}'
    def request(path, data=None, token=True, headers=None):
        hdr=dict(headers or {})
        if token:hdr['X-Studio-Token']=server.app.token
        with urllib.request.urlopen(urllib.request.Request(base+path,data=data,headers=hdr),timeout=10) as response:
            return json.load(response)
    try:
        with pytest.raises(urllib.error.HTTPError) as exc:
            request('/api/save',b'{}',token=False)
        assert exc.value.code==403
        image, region=sample(fonts);buf=io.BytesIO();Image.fromarray(image).save(buf,format='PNG')
        project=request('/api/import',buf.getvalue(),headers={'X-Filename':'sample.png'})
        assert project['pages'][0]['width']==700
        project['pages'][0]['regions']=[region]
        saved=request('/api/save',json.dumps(project).encode())
        assert saved['revision']==project['revision']+1
        assert saved['pages'][0]['render_revision']==-1
        with pytest.raises(urllib.error.HTTPError) as exc:
            request('/api/save',json.dumps(project).encode())
        assert exc.value.code==409
        with pytest.raises(urllib.error.HTTPError):
            request('/asset?project='+project['id']+'&file=../../etc/passwd')
    finally:
        server.shutdown();server.server_close();thread.join();server.app.executor.shutdown()


def batch_project(tmp_path, fonts):
    from ai_text_sharpener.studio import Studio
    app = Studio(tmp_path)
    image, region = sample(fonts)
    project = app.store.create("Batch")
    app.store.add_images(project, [(str(i), Image.fromarray(image)) for i in range(3)])
    return app, project, region


def run_background(app, project, kind, **extra):
    app.start_job({"kind": kind, "project_id": project["id"], "revision": project["revision"], **extra})
    # Wait on the same single-worker queue, with a deadline to catch deadlocks.
    app.executor.submit(lambda: None).result(timeout=20)
    return app.store.load(project["id"])


def test_batch_preserves_edits_and_exports_every_page(tmp_path, fonts, monkeypatch):
    from copy import deepcopy
    from ai_text_sharpener import studio
    from pptx import Presentation
    app, project, region = batch_project(tmp_path, fonts)
    project["pages"][0]["regions"] = [dict(region, x=55, locked=True)]
    app.store.save(project)
    calls = []
    def detect(path, progress):
        calls.append(path.stem)
        return [deepcopy(region)]
    monkeypatch.setattr(studio, "detect_regions", detect)
    monkeypatch.setattr(studio, "fit_region", lambda image, r, *args: r)
    try:
        saved = run_background(app, project, "analyze_all")
        assert app.job["status"] == "done"
        assert (app.job["completed"], app.job["skipped"], app.job["failed"]) == (2,1,0)
        assert saved["pages"][0]["regions"][0]["x"] == 55
        assert len(calls) == 2
        saved = run_background(app, saved, "analyze_all")
        assert app.job["skipped"] == 3 and len(calls) == 2
        saved = run_background(app, saved, "export")
        assert app.job["status"] == "done" and app.job["download"]
        path = app.store.directory(saved["id"]) / "export.pptx"
        deck = Presentation(path)
        assert len(deck.slides) == 3
        with zipfile.ZipFile(path) as archive:
            assert all(f'ppt/media/image_vector_{i}.svg' in archive.namelist() for i in range(1,4))
        assert all(p["render_revision"] == saved["revision"] for p in saved["pages"])
    finally:
        app.executor.shutdown()


def test_batch_failure_keeps_order_and_retries_only_failed_page(tmp_path, fonts, monkeypatch):
    from ai_text_sharpener import studio
    app, project, _ = batch_project(tmp_path, fonts)
    bad_id = project["pages"][1]["id"]
    calls = []
    def detect(path, progress):
        calls.append(path.stem)
        if path.stem == bad_id: raise RuntimeError("test OCR failure")
        return []
    monkeypatch.setattr(studio, "detect_regions", detect)
    try:
        saved = run_background(app, project, "analyze_all")
        assert app.job["failed"] == 1 and app.job["completed"] == 2
        assert [p["status"] for p in saved["pages"]] == ["review", "error", "review"]
        assert [p["id"] for p in saved["pages"]] == [p["id"] for p in project["pages"]]
        monkeypatch.setattr(studio, "detect_regions", lambda path, progress: [])
        saved = run_background(app, saved, "analyze_all")
        assert app.job["completed"] == 1 and app.job["skipped"] == 2
        assert "error" not in saved["pages"][1]
    finally:
        app.executor.shutdown()


def test_batch_cancel_checkpoints_completed_pages(tmp_path, fonts, monkeypatch):
    from ai_text_sharpener import studio
    app, project, _ = batch_project(tmp_path, fonts)
    count = 0
    def detect(path, progress):
        nonlocal count
        count += 1
        if count == 2:
            app.cancelled.set()
            progress("cancel now")
        return []
    monkeypatch.setattr(studio, "detect_regions", detect)
    try:
        saved = run_background(app, project, "analyze_all")
        assert app.job["status"] == "cancelled"
        assert [p["status"] for p in saved["pages"]] == ["review", "new", "new"]
        monkeypatch.setattr(studio, "detect_regions", lambda path, progress: [])
        saved = run_background(app, saved, "analyze_all")
        assert app.job["completed"] == 2 and app.job["skipped"] == 1
    finally:
        app.executor.shutdown()


def test_full_rescan_replaces_completed_and_locked_regions_and_backs_up(tmp_path, fonts, monkeypatch):
    from copy import deepcopy
    from ai_text_sharpener import studio
    app, project, region = batch_project(tmp_path, fonts)
    for p in project['pages']:
        p.update(status='review', regions=[dict(region, x=55, locked=True)])
    app.store.save(project)
    calls=[]
    def detect(path, progress):
        calls.append(path.stem)
        return [deepcopy(region)]
    monkeypatch.setattr(studio,'detect_regions',detect)
    monkeypatch.setattr(studio,'fit_region',lambda image,r,*args:r)
    try:
        saved=run_background(app,project,'analyze_all',reset=True)
        assert app.job['completed']==3 and app.job['skipped']==0
        assert calls==[p['id'] for p in project['pages']]
        assert all(p['regions'][0]['x']==40 and not p['regions'][0]['locked'] for p in saved['pages'])
        assert all(not p.get('pending_rescan') for p in saved['pages'])
        backups=list((app.store.directory(project['id'])/'backups').glob('before-rescan-*.json'))
        assert len(backups)==1
        assert json.loads(backups[0].read_text())['pages'][0]['regions'][0]['x']==55
    finally:app.executor.shutdown()


@pytest.mark.parametrize('failure', ['cancel','error'])
def test_full_rescan_retains_unfinished_results_and_resumes(tmp_path, fonts, monkeypatch, failure):
    from copy import deepcopy
    from ai_text_sharpener import studio
    app, project, region = batch_project(tmp_path, fonts)
    for p in project['pages']:
        p.update(status='review',regions=[dict(region,x=55,locked=True)])
    app.store.save(project)
    count=0
    def detect(path,progress):
        nonlocal count
        count+=1
        if count==2:
            if failure=='cancel':
                app.cancelled.set();progress('cancel')
            else:raise ValueError('test OCR failure')
        return [deepcopy(region)]
    monkeypatch.setattr(studio,'detect_regions',detect)
    monkeypatch.setattr(studio,'fit_region',lambda image,r,*args:r)
    try:
        saved=run_background(app,project,'analyze_all',reset=True)
        assert saved['pages'][1]['pending_rescan']
        assert saved['pages'][1]['regions'][0]['x']==55
        assert not saved['pages'][0].get('pending_rescan')
        monkeypatch.setattr(studio,'detect_regions',lambda path,progress:[deepcopy(region)])
        resumed=run_background(app,saved,'analyze_all')
        assert app.job['completed']==(2 if failure=='cancel' else 1)
        assert all(p['regions'][0]['x']==40 for p in resumed['pages'])
        assert all(not p.get('pending_rescan') for p in resumed['pages'])
    finally:app.executor.shutdown()


@pytest.mark.parametrize('scope', ['current', 'all'])
@pytest.mark.parametrize('format', ['pptx', 'svg', 'png'])
def test_export_scope_format_and_download(tmp_path, fonts, scope, format):
    """Both scopes export current edits, in order, through the real asset route."""
    from copy import deepcopy
    from urllib.parse import parse_qs, urlparse
    from pptx import Presentation
    from ai_text_sharpener.studio_project import compose_page
    server = create_server(tmp_path, 0)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    app = server.app
    image, region = sample(fonts)
    project = app.store.create('Export')
    app.store.add_images(project, [('same/name', Image.fromarray(image)) for _ in range(3)])
    for index, item in enumerate(project['pages']):
        item['regions'] = [dict(deepcopy(region), x=40+index*15, text=f'Edited {index}')]
    app.store.save(project)
    directory = app.store.directory(project['id'])
    selected = project['pages'][1]
    included = project['pages'] if scope == 'all' else [selected]
    expected = [compose_page(directory, deepcopy(p)).encode() for p in included]
    try:
        saved = run_background(app, project, 'export', scope=scope, format=format, page_id=selected['id'])
        assert app.job['status'] == 'done', app.job
        url = f'http://127.0.0.1:{server.server_port}' + app.job['download']
        with urllib.request.urlopen(url) as response:
            content = response.read()
            assert 'attachment' in response.headers['Content-Disposition']
        name = parse_qs(urlparse(url).query)['file'][0]
        assert len(saved['pages']) == 3
        for item in saved['pages']:
            assert item['render_revision'] == (saved['revision'] if scope == 'all' or item['id'] == selected['id'] else -1)
        if format == 'pptx':
            deck = Presentation(io.BytesIO(content))
            assert len(deck.slides) == len(included)
            assert name == ('export.pptx' if scope == 'all' else 'export-page.pptx')
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                for index, svg in enumerate(expected, 1):
                    assert archive.read(f'ppt/media/image_vector_{index}.svg') == svg
        else:
            if scope == 'all':
                assert name == f'export-{format}.zip'
                with zipfile.ZipFile(io.BytesIO(content)) as archive:
                    assert archive.namelist() == [f'{i:03d}.{format}' for i in range(1, 4)]
                    results = [archive.read(n) for n in archive.namelist()]
            else:
                assert name == f"{selected['id']}-result.{format}"
                results = [content]
            for result, item, svg in zip(results, included, expected):
                if format == 'svg':
                    assert result == svg
                else:
                    import resvg_py
                    assert np.array_equal(np.asarray(Image.open(io.BytesIO(result))),
                                          np.asarray(Image.open(io.BytesIO(resvg_py.svg_to_bytes(svg_string=svg.decode())))))
                    assert Image.open(io.BytesIO(result)).size == (item['width'], item['height'])
    finally:
        server.shutdown();server.server_close();thread.join();app.executor.shutdown()


@pytest.mark.parametrize('extra', [{'scope':'missing'}, {'format':'pdf'}, {'scope':'current','page_id':'missing'}])
def test_invalid_export_is_rejected_before_job(tmp_path, fonts, extra):
    app, project, _ = batch_project(tmp_path, fonts)
    try:
        with pytest.raises(ValueError):
            app.start_job(dict(kind='export', project_id=project['id'], revision=project['revision'], **extra))
        assert not app.job['active']
        assert not list(app.store.directory(project['id']).glob('export*'))
    finally:
        app.executor.shutdown()


def test_cancelled_image_export_keeps_previous_archive(tmp_path, fonts):
    from ai_text_sharpener.studio_project import export_pages
    app, project, _ = batch_project(tmp_path, fonts)
    directory = app.store.directory(project['id'])
    archive = directory / 'export-svg.zip'
    archive.write_bytes(b'previous export')
    def cancel(message, value):
        if value > 0:
            raise InterruptedError('cancelled')
    try:
        with pytest.raises(InterruptedError):
            export_pages(directory, project, format='svg', progress=cancel)
        assert archive.read_bytes() == b'previous export'
        assert not (directory / 'export-svg.tmp').exists()
    finally:
        app.executor.shutdown()
