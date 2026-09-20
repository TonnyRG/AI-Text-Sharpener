"""Whole-slide raster import tolerates subpixel fitting, not visible layout changes."""
import io

import pytest
from PIL import Image
from pptx import Presentation

from ai_text_sharpener.studio_project import extract_image_slides


def picture_deck(size=(1672, 941), edges=(3239, 0, 3239, 0), transform=None):
    deck = Presentation()
    deck.slide_width, deck.slide_height = 12192000, 6858000
    image = Image.new('RGB', size, '#17496d')
    image.putpixel((0, 0), (255, 10, 20))
    image.putpixel((size[0]-1, size[1]-1), (20, 240, 10))
    blob = io.BytesIO()
    image.save(blob, format='PNG')
    left, top, right, bottom = edges
    for _ in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_picture(io.BytesIO(blob.getvalue()), left, top,
            deck.slide_width-left-right, deck.slide_height-top-bottom)
        if transform:
            transform(shape)
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue(), image


def test_centered_aspect_fit_import_preserves_every_source_pixel_and_page_order():
    # Same geometry as the reported 16:9 deck: 0.444 native pixels per side.
    data, original = picture_deck()
    pages = extract_image_slides(data)
    assert [name for name, _ in pages] == ['第 1 页', '第 2 页']
    for _, image in pages:
        assert image.size == original.size
        assert image.tobytes() == original.tobytes()


@pytest.mark.parametrize('edges', [(6000, 6000, 6000, 6000), (-6000, -6000, 6000, 6000)])
def test_subpixel_rounding_allowed_on_all_four_edges(edges):
    data, original = picture_deck(edges=edges)
    assert extract_image_slides(data)[0][1].tobytes() == original.tobytes()


@pytest.mark.parametrize('size,edges', [
    ((1672, 941), (15000, 0, 0, 0)),  # visible left margin
    ((1672, 941), (0, 0, 15000, 0)),  # only right edge differs
    ((1672, 941), (0, 15000, 0, 0)),
    ((1672, 941), (0, 0, 0, 15000)),
    ((1672, 941), (-15000, 0, 15000, 0)),  # picture clipped by slide
    ((3840, 2160), (6000, 0, 0, 0)),  # tolerance follows source resolution
])
def test_visible_borders_and_offsets_still_require_rendered_input(size, edges):
    data, _ = picture_deck(size=size, edges=edges)
    with pytest.raises(ValueError, match='边缘偏差超过 1 像素'):
        extract_image_slides(data)


def flip(shape):
    shape._element.xpath('.//a:xfrm')[0].set('flipH', '1')


@pytest.mark.parametrize('transform', [
    lambda s: setattr(s, 'rotation', 5),
    lambda s: setattr(s, 'crop_left', .05),
    flip,
])
def test_tolerance_does_not_bypass_cropping_rotation_or_flip(transform):
    data, _ = picture_deck(transform=transform)
    with pytest.raises(ValueError, match='裁剪、旋转'):
        extract_image_slides(data)
