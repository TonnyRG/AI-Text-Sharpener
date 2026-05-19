"""Tests for the flat-image PPTX exporter."""
from pathlib import Path

import pytest
from PIL import Image

from ai_text_sharpener.ppt_export import export_project_to_pptx
from ai_text_sharpener.project import (
    ReviewProject,
    ReviewProjectItem,
    write_review_project,
)


def _make_png(path: Path, size=(160, 90), color="white") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, color).save(path)


def _make_project(tmp_path: Path, *, with_final: bool, with_draft: bool) -> Path:
    items = []
    for i in range(2):
        stem = f"slide_{i+1:03d}"
        final = tmp_path / f"{stem}_final.png"
        draft = tmp_path / f"{stem}_draft.png"
        if with_final:
            _make_png(final, color="red" if i == 0 else "blue")
        if with_draft:
            _make_png(draft, color="green")
        items.append(ReviewProjectItem(
            id=f"slide-{i+1:03d}",
            name=stem,
            image_path=f"{stem}.jpg",
            review_path=f"{stem}_review.json",
            output_png=final.name if with_final else "",
            draft_png=draft.name if with_draft else "",
        ))
    project_path = tmp_path / "review_project.json"
    write_review_project(ReviewProject(items=items), project_path)
    return project_path


def test_export_uses_final_png(tmp_path):
    project_path = _make_project(tmp_path, with_final=True, with_draft=False)
    output = tmp_path / "deck.pptx"
    result = export_project_to_pptx(project_path, output)
    assert result == output
    assert output.exists()
    assert output.stat().st_size > 1000

    from pptx import Presentation
    prs = Presentation(str(output))
    assert len(prs.slides) == 2


def test_export_falls_back_to_draft_when_final_missing(tmp_path):
    project_path = _make_project(tmp_path, with_final=False, with_draft=True)
    output = tmp_path / "deck.pptx"
    result = export_project_to_pptx(project_path, output)
    assert output.exists()

    from pptx import Presentation
    prs = Presentation(str(output))
    assert len(prs.slides) == 2


def test_export_errors_when_no_rendered_images(tmp_path):
    project_path = _make_project(tmp_path, with_final=False, with_draft=False)
    with pytest.raises(RuntimeError, match="No rendered PNG"):
        export_project_to_pptx(project_path, tmp_path / "deck.pptx")


def test_export_slide_aspect_matches_first_image(tmp_path):
    items = [
        ReviewProjectItem(
            id="slide-001", name="s1",
            image_path="s1.jpg", review_path="s1.json",
            output_png="s1_final.png",
        ),
    ]
    _make_png(tmp_path / "s1_final.png", size=(2000, 1000))
    project_path = tmp_path / "p.json"
    write_review_project(ReviewProject(items=items), project_path)

    output = tmp_path / "d.pptx"
    export_project_to_pptx(project_path, output)

    from pptx import Presentation
    prs = Presentation(str(output))
    assert abs(prs.slide_width / prs.slide_height - 2.0) < 0.01
