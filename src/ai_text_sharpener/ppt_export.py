"""Export a review project to a flat-image PPTX (one rendered slide per page)."""
from __future__ import annotations

import argparse
import tempfile
from pathlib import Path

from PIL import Image

from .project import (
    ReviewProjectItem,
    load_review_project,
    resolve_project_path,
)

EMU_PER_INCH = 914400
DEFAULT_SLIDE_WIDTH_INCHES = 13.333  # standard PPT 16:9 width
DEFAULT_MAX_WIDTH_PX = 3840           # 2x 1080p; matches a 4K projector worst-case
DEFAULT_JPEG_QUALITY = 92             # visually lossless for slide content


def _resolve_render_path(project_path: Path, item: ReviewProjectItem) -> Path:
    """Return the best available rendered PNG: final, falling back to draft."""
    for value in (item.output_png, item.draft_png):
        if not value:
            continue
        candidate = resolve_project_path(project_path, value)
        if candidate.exists():
            return candidate
    raise RuntimeError(
        f"No rendered PNG for slide {item.id} ({item.name}). "
        "Click Render in the editor before exporting."
    )


def _shrink_to_jpeg(src: Path, dst: Path, max_width_px: int, jpeg_quality: int) -> None:
    """Resize ``src`` (PNG/RGBA) and save as flattened JPEG ``dst``.

    PNGs from the pipeline are 8000×4500 RGBA — far above any display resolution
    PPT will use.  Downsampling to ``max_width_px`` and re-encoding as JPEG cuts
    file size ~10–20× with no visible loss (slides composite onto opaque
    backgrounds, so alpha and lossless storage are both wasted).
    """
    with Image.open(src) as im:
        if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
            bg = Image.new("RGB", im.size, (255, 255, 255))
            bg.paste(im, mask=im.convert("RGBA").split()[-1])
            im = bg
        elif im.mode != "RGB":
            im = im.convert("RGB")
        w, h = im.size
        if w > max_width_px:
            new_h = int(round(h * max_width_px / w))
            im = im.resize((max_width_px, new_h), Image.LANCZOS)
        im.save(dst, format="JPEG", quality=jpeg_quality, optimize=True, progressive=True)


def export_project_to_pptx(
    project_path: Path,
    output_pptx: Path,
    slide_width_inches: float = DEFAULT_SLIDE_WIDTH_INCHES,
    max_width_px: int = DEFAULT_MAX_WIDTH_PX,
    jpeg_quality: int = DEFAULT_JPEG_QUALITY,
) -> Path:
    """Write a flat-image PPTX where each slide is the rendered image full-bleed.

    Slide size = ``slide_width_inches`` × (height derived from first image's
    aspect ratio).  PPT decks require a single slide size for all pages, so the
    first slide's aspect ratio wins.  Standard PPT 16:9 width is 13.333 inches.

    Pictures are downsampled to ``max_width_px`` and re-encoded as JPEG at
    ``jpeg_quality`` before embedding, keeping the deck small while remaining
    visually lossless at projector resolutions.  Pass ``max_width_px=0`` to
    disable resizing.
    """
    from pptx import Presentation

    project_path = Path(project_path)
    output_pptx = Path(output_pptx)
    project = load_review_project(project_path)
    if not project.items:
        raise RuntimeError(f"Project has no items: {project_path}")

    images = [(item, _resolve_render_path(project_path, item)) for item in project.items]

    with Image.open(images[0][1]) as im:
        first_w, first_h = im.size
    slide_w_emu = int(slide_width_inches * EMU_PER_INCH)
    slide_h_emu = int(slide_w_emu * first_h / first_w)

    prs = Presentation()
    prs.slide_width = slide_w_emu
    prs.slide_height = slide_h_emu
    blank_layout = prs.slide_layouts[6]

    effective_max = max_width_px if max_width_px and max_width_px > 0 else None

    with tempfile.TemporaryDirectory(prefix="ats_pptx_") as tmpdir:
        tmp = Path(tmpdir)
        for idx, (_, png) in enumerate(images):
            picture_path = png
            if effective_max is not None:
                jpg_path = tmp / f"slide_{idx:03d}.jpg"
                _shrink_to_jpeg(png, jpg_path, effective_max, jpeg_quality)
                picture_path = jpg_path
            slide = prs.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                str(picture_path), 0, 0,
                width=prs.slide_width, height=prs.slide_height,
            )

        output_pptx.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output_pptx))
    return output_pptx


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Export a review project to a flat-image PPTX."
    )
    parser.add_argument("project_path", type=Path)
    parser.add_argument("--output", "-o", type=Path, default=None)
    parser.add_argument(
        "--width-inches", type=float, default=DEFAULT_SLIDE_WIDTH_INCHES,
        help="Slide width in inches (default 13.333 = standard 16:9).",
    )
    parser.add_argument(
        "--max-width-px", type=int, default=DEFAULT_MAX_WIDTH_PX,
        help=(
            "Downsample each picture to this width before embedding "
            f"(default {DEFAULT_MAX_WIDTH_PX}; pass 0 to keep original size)."
        ),
    )
    parser.add_argument(
        "--jpeg-quality", type=int, default=DEFAULT_JPEG_QUALITY,
        help=f"JPEG quality 1-100 (default {DEFAULT_JPEG_QUALITY}).",
    )
    args = parser.parse_args()

    if not args.project_path.exists():
        raise SystemExit(f"Project file not found: {args.project_path}")

    output = args.output or args.project_path.parent / (
        args.project_path.stem + "_export.pptx"
    )
    result = export_project_to_pptx(
        args.project_path, output,
        slide_width_inches=args.width_inches,
        max_width_px=args.max_width_px,
        jpeg_quality=args.jpeg_quality,
    )
    print(f"PPTX -> {result}")


if __name__ == "__main__":
    main()
