"""Portable Studio projects, input adapters and vector-preserving outputs."""
from __future__ import annotations

import base64
import io
import json
import math
import os
import re
import tempfile
import time
import uuid
import zipfile
from html import escape
from pathlib import Path

import numpy as np
import resvg_py
from lxml import etree
from PIL import Image, ImageOps

from .fidelity import erase_regions, automatic_fit_failure
from .formula import formula_layout, overlaps
from .colors import region_layout, validate_colors
from .typography import font_catalog, svg_document

MAX_PIXELS = 40_000_000


def data_root() -> Path:
    return Path(os.environ.get("XDG_DATA_HOME", str(Path.home() / ".local/share"))) / "ai-text-sharpener"


class StudioStore:
    def __init__(self, root=None):
        self.root = Path(root) if root else data_root()
        (self.root / "projects").mkdir(parents=True, exist_ok=True)

    def directory(self, project_id):
        if not re.fullmatch(r"[0-9a-f]{32}", str(project_id)):
            raise ValueError("无效项目编号。")
        return self.root / "projects" / project_id

    def load(self, project_id):
        return json.loads((self.directory(project_id) / "project.json").read_text())

    def save(self, project):
        project["updated"] = time.time()
        target = self.directory(project["id"]) / "project.json"
        target.parent.mkdir(exist_ok=True)
        # Atomic replacement prevents partial JSON on interruption.
        fd, name = tempfile.mkstemp(dir=target.parent, suffix=".tmp")
        try:
            with os.fdopen(fd, "w") as handle:
                json.dump(project, handle, ensure_ascii=False, indent=2)
            os.replace(name, target)
        finally:
            Path(name).unlink(missing_ok=True)

    def listing(self, trashed=False):
        items = []
        for path in (self.root / ("trash" if trashed else "projects")).glob("*/project.json"):
            try:
                doc = json.loads(path.read_text())
                items.append({"id": doc["id"], "name": doc["name"],
                              "updated": doc["updated"], "revision": doc["revision"], "pages": len(doc["pages"])})
            except (ValueError, KeyError, OSError):
                continue
        return sorted(items, key=lambda p: p["updated"], reverse=True)

    def trash(self, project_id, revision):
        project = self.load(project_id)
        if revision != project["revision"]:
            raise RuntimeError("工作区已发生变化，请刷新管理列表后重试。")
        target = self.root / "trash" / project_id
        target.parent.mkdir(exist_ok=True)
        if target.exists():
            raise RuntimeError("回收站已有同编号工作区。")
        self.directory(project_id).rename(target)

    def restore(self, project_id):
        target = self.directory(project_id)  # Validate ID before using it as a path.
        if target.exists():
            raise RuntimeError("工作区已存在，请刷新管理列表。")
        (self.root / "trash" / project_id).rename(target)

    def create(self, name):
        result = {"version": 2, "id": uuid.uuid4().hex, "name": str(name)[:160],
                  "revision": 0, "pages": [], "updated": time.time()}
        self.save(result)
        return result

    def add_images(self, project, images):
        for name, image in images:
            image = ImageOps.exif_transpose(image).convert("RGBA")
            if image.width * image.height > MAX_PIXELS:
                raise ValueError("单张图片不得超过 4000 万像素，请先缩小图片。")
            flat = Image.new("RGBA", image.size, "white")
            flat.alpha_composite(image)
            page_id = uuid.uuid4().hex[:12]
            flat.convert("RGB").save(self.directory(project["id"]) / f"{page_id}.png")
            project["pages"].append({"id": page_id, "name": name[:160],
                                     "width": image.width, "height": image.height,
                                     "regions": [], "status": "new", "render_revision": -1})
        project["revision"] += 1
        self.save(project)
        return project


def image_inputs(data: bytes, filename: str):
    if filename.lower().endswith(".pptx"):
        return extract_image_slides(data)
    with Image.open(io.BytesIO(data)) as im:
        if im.width * im.height > MAX_PIXELS:
            raise ValueError("图片超过 4000 万像素。")
        im.load()
        return [(Path(filename).stem, im.copy())]


def extract_image_slides(data: bytes):
    """Lossless extraction only for provably single full-bleed raster slides.

    Refuse composed decks instead of silently losing native text or drawings.
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        if sum(info.file_size for info in package.infolist()) > 500_000_000:
            raise ValueError("PPTX 解压后过大。")
    deck = Presentation(io.BytesIO(data))
    results = []
    for number, slide in enumerate(deck.slides, 1):
        if len(slide.shapes) != 1 or slide.shapes[0].shape_type != MSO_SHAPE_TYPE.PICTURE:
            raise ValueError(f"第 {number} 页不是单张铺满图片。请将复杂 PPT 导出为 PNG/JPG 后导入。")
        shape = slide.shapes[0]
        if (abs(shape.left) > 100 or abs(shape.top) > 100 or
            abs(shape.width - deck.slide_width) > 100 or
            abs(shape.height - deck.slide_height) > 100 or shape.rotation or
            any(abs(v) > .00001 for v in (shape.crop_left, shape.crop_right, shape.crop_top, shape.crop_bottom)) or
            shape._element.xpath('.//a:xfrm[@flipH="1" or @flipV="1"]')):
            raise ValueError(f"第 {number} 页包含裁剪、旋转或非铺满图片，请先导出该页为 PNG。")
        # SVG pictures carry a raster fallback: extracting just that fallback
        # would lose precision, so ask for an explicitly rendered input.
        if shape._element.xpath('.//*[local-name()="svgBlip"]'):
            raise ValueError(f"第 {number} 页是矢量图片，请使用原始 PNG/JPG 输入。")
        with Image.open(io.BytesIO(shape.image.blob)) as image:
            if image.width * image.height > MAX_PIXELS:
                raise ValueError(f"第 {number} 页图片过大。")
            image.load()
            if abs(image.width / image.height - deck.slide_width / deck.slide_height) > .015:
                raise ValueError(f"第 {number} 页图片被拉伸，请先导出为 PNG。")
            results.append((f"第 {number} 页", image.copy()))
    if not results:
        raise ValueError("PPTX 没有页面。")
    return results


def validate_regions(regions, width, height):
    if not isinstance(regions, list) or len(regions) > 500:
        raise ValueError("最多支持 500 个文字区域。")
    ids = set()
    for region in regions:
        if not isinstance(region.get("id"), str) or region["id"] in ids:
            raise ValueError("文字区域编号无效或重复。")
        ids.add(region["id"])
        if region.get("kind", "text") not in {"text", "formula"}:
            raise ValueError("未知区域类型。")
        if region.get("kind") == "formula":
            if not isinstance(region.get("latex", ""), str) or len(region.get("latex", "")) > 2048:
                raise ValueError("公式 LaTeX 最多 2048 个字符。")
        if not isinstance(region.get("text"), str) or len(region["text"]) > 512 or "\n" in region["text"]:
            raise ValueError("每个区域支持最多 512 个字符的单行文字。")
        if len(region.get("bbox", [])) != 4:
            raise ValueError("无效文字框。")
        values = region["bbox"] + [region.get("x"), region.get("y"), region.get("font_size"), region.get("letter_spacing", 0), region.get("stroke_width", 0)]
        if any(not isinstance(v, (int, float)) or not math.isfinite(v) for v in values):
            raise ValueError("坐标和字号必须是有限数值。")
        bx, by, bw, bh = region["bbox"]
        if bw < 3 or bh < 3 or bw > width * 2 or bh > height * 2:
            raise ValueError("文字框尺寸无效。")
        if not (-width <= region["x"] <= width * 2 and -height <= region["y"] <= height * 2):
            raise ValueError("文字位置超出允许范围。")
        if not 4 <= region["font_size"] <= 1000 or abs(region.get("letter_spacing", 0)) > 300:
            raise ValueError("字号或字距超出允许范围。")
        if not 0 <= region.get("stroke_width", 0) <= 12:
            raise ValueError("笔画加粗必须在 0–12 像素之间。")
        if region.get("kind") != "formula" and region.get("enabled", True) and region.get("font_id") not in font_catalog():
            raise ValueError("所选字体在本机不可用，请重新选择。")
        if not re.fullmatch(r"#[0-9a-fA-F]{6}", region.get("color", "")):
            raise ValueError("无效文字颜色。")
        validate_colors(region)
        if region.get("erase_mode", "gradient") not in {"gradient", "inpaint", "none"}:
            raise ValueError("无效背景修复方式。")


def compose_page(directory: Path, page: dict):
    """Build the exact export SVG in memory; preview must not write project files."""
    source = np.asarray(Image.open(directory / f"{page['id']}.png").convert("RGB"))
    validate_regions(page["regions"], page["width"], page["height"])
    formulas = [r for r in page["regions"] if r.get("kind") == "formula"]
    enabled, layouts = [], {}
    for region in page["regions"]:
        if not region.get("enabled", True):
            continue
        if region.get("kind") == "formula":
            layout = formula_layout(region.get("latex", ""), region["font_size"], region.get("stroke_width", 0))
        else:
            # Old/manual OCR fragments must not erase or paint over a formula.
            if any(overlaps(region["bbox"], f["bbox"]) for f in formulas):
                continue
            if not region["text"].strip():
                enabled.append(region)
                continue
            layout = region_layout(region)
            if automatic_fit_failure(region, layout):
                continue
        layouts[region["id"]] = layout
        enabled.append(region)
    background = erase_regions(source, enabled)
    buf = io.BytesIO()
    Image.fromarray(background).save(buf, format="PNG")
    body = [f'<image width="{page["width"]}" height="{page["height"]}" xlink:href="data:image/png;base64,{base64.b64encode(buf.getvalue()).decode()}"/>']
    for region in enabled:
        if region["id"] not in layouts:
            region["ink_width"] = region["ink_height"] = 0
            continue
        layout = layouts[region["id"]]
        region["ink_width"] = round(layout["width"], 2)
        region["ink_height"] = round(layout["height"], 2)
        body.append(f'<g data-region-id="{escape(region["id"], quote=True)}" transform="translate({region["x"]} {region["y"]})" fill="{region["color"]}" color="{region["color"]}">{layout["body"]}</g>')
    # Explicit preservation restores the complete source crop last, even when
    # a neighboring erase mask or manually moved text overlaps it.
    for region in page["regions"]:
        if region.get("enabled") or not (region.get("kind") == "formula" or region.get("preserve_original")):
            continue
        x, y, w, h = region["bbox"]
        x0,y0=max(0,int(x)),max(0,int(y));x1,y1=min(source.shape[1],math.ceil(x+w)),min(source.shape[0],math.ceil(y+h))
        if x1 <= x0 or y1 <= y0:
            continue
        buffer=io.BytesIO();Image.fromarray(source[y0:y1,x0:x1]).save(buffer,format="PNG")
        body.append(f'<image x="{x0}" y="{y0}" width="{x1-x0}" height="{y1-y0}" xlink:href="data:image/png;base64,{base64.b64encode(buffer.getvalue()).decode()}"/>')
    return svg_document(page["width"], page["height"], "".join(body))


def render_page(directory: Path, page: dict):
    svg = compose_page(directory, page)
    png = bytes(resvg_py.svg_to_bytes(svg_string=svg))
    # Publish both after successful rendering; no half-written output files.
    for suffix, content in [("svg", svg.encode()), ("png", png)]:
        path = directory / f"{page['id']}-result.{suffix}"
        temp = path.with_suffix(path.suffix + ".tmp")
        temp.write_bytes(content)
        os.replace(temp, path)
    return svg, png


def export_pages(directory: Path, project: dict, *, scope="all", format="pptx", page_id=None, progress=None) -> Path:
    """Export the requested pages using the same renderer as the live preview."""
    if scope not in {"all", "current"} or format not in {"pptx", "svg", "png", "json"}:
        raise ValueError("无效的导出范围或格式。")
    pages = project["pages"] if scope == "all" else [p for p in project["pages"] if p["id"] == page_id]
    if not pages:
        raise ValueError("没有可导出的页面。")
    if format == "json":
        if scope != "all":
            raise ValueError("编辑数据备份必须包含全部页面。")
        return directory / "project.json"
    if format == "pptx":
        return export_vector_pptx(directory, dict(project, pages=pages), progress,
                                  filename="export.pptx" if scope == "all" else "export-page.pptx")
    if scope == "current":
        if progress:
            progress("正在导出当前页…", 10)
        render_page(directory, pages[0])
        return directory / f"{pages[0]['id']}-result.{format}"
    target = directory / f"export-{format}.zip"
    temp = target.with_suffix(".tmp")
    try:
        with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as archive:
            for index, page in enumerate(pages, 1):
                if progress:
                    progress(f"导出第 {index}/{len(pages)} 页", (index-1) / len(pages) * 95)
                svg, png = render_page(directory, page)
                archive.writestr(f"{index:03d}.{format}", svg.encode() if format == "svg" else png)
        if progress:
            progress("正在打包 ZIP…", 98)
        os.replace(temp, target)
    finally:
        temp.unlink(missing_ok=True)
    return target


def export_vector_pptx(directory: Path, project: dict, progress=None, *, filename="export.pptx") -> Path:
    """Office SVG picture extension plus lossless PNG fallback.

    Text is outlined, not editable PowerPoint text. Letterboxing preserves
    aspect ratio for mixed-size pages. All outputs are re-rendered first.
    """
    from pptx import Presentation
    from pptx.util import Inches
    if not project["pages"]:
        raise ValueError("项目没有页面。")
    presentation = Presentation()
    first = project["pages"][0]
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = int(presentation.slide_width * first["height"] / first["width"])
    svgs = []
    for index, page in enumerate(project["pages"]):
        if progress:
            progress(f"导出第 {index+1}/{len(project['pages'])} 页", index / len(project["pages"]) * 95)
        svg, png = render_page(directory, page)
        scale = min(presentation.slide_width / page["width"], presentation.slide_height / page["height"])
        w, h = int(page["width"] * scale), int(page["height"] * scale)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(png), (presentation.slide_width - w) // 2,
                                 (presentation.slide_height - h) // 2, width=w, height=h)
        slide.notes_slide.notes_text_frame.text = "AI Text Sharpener · 保真文字轮廓（在 Studio 中编辑）\n" + "\n".join(r.get("latex", "") if r.get("kind") == "formula" else r["text"] for r in page["regions"])
        svgs.append(svg.encode())
    raw = io.BytesIO()
    presentation.save(raw)
    with zipfile.ZipFile(raw) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    rns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ans = "http://schemas.openxmlformats.org/drawingml/2006/main"
    ons = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    sns = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
    for idx, svg in enumerate(svgs, 1):
        media = f"image_vector_{idx}.svg"
        parts[f"ppt/media/{media}"] = svg
        relpath = f"ppt/slides/_rels/slide{idx}.xml.rels"
        rels = etree.fromstring(parts[relpath])
        used = {r.attrib["Id"] for r in rels}
        n = 1
        while f"rId{n}" in used:
            n += 1
        rid = f"rId{n}"
        etree.SubElement(rels, f"{{{rns}}}Relationship", Id=rid, Type=ons + "/image", Target=f"../media/{media}")
        parts[relpath] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
        slidepath = f"ppt/slides/slide{idx}.xml"
        slide = etree.fromstring(parts[slidepath])
        blip = slide.find(f".//{{{ans}}}blip")
        extlist = etree.SubElement(blip, f"{{{ans}}}extLst")
        ext = etree.SubElement(extlist, f"{{{ans}}}ext", uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
        etree.SubElement(ext, f"{{{sns}}}svgBlip", {f"{{{ons}}}embed": rid}, nsmap={"asvg": sns})
        parts[slidepath] = etree.tostring(slide, xml_declaration=True, encoding="UTF-8", standalone=True)
    types = etree.fromstring(parts["[Content_Types].xml"])
    tns = "http://schemas.openxmlformats.org/package/2006/content-types"
    etree.SubElement(types, f"{{{tns}}}Default", Extension="svg", ContentType="image/svg+xml")
    parts["[Content_Types].xml"] = etree.tostring(types, xml_declaration=True, encoding="UTF-8", standalone=True)
    if progress:
        progress("正在打包 PPTX…", 98)
    target = directory / filename
    temp = target.with_suffix(".tmp")
    with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    os.replace(temp, target)
    return target
